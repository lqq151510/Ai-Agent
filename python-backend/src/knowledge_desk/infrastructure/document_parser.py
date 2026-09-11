"""Native document parsing for the import pipeline.

Covers the formats the product declares: PDF, DOCX, PPTX, Markdown, TXT and
HTML. Format metadata is retained so it can be surfaced on the knowledge item
without exposing anything about the original host path.

Native parsers are the primary path (deterministic and dependency-light).
MarkItDown is used only as an optional fallback when it happens to be
installed and a native parser cannot handle the payload.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from html.parser import HTMLParser
from typing import Any

from knowledge_desk.domain.enums import SourceType
from knowledge_desk.errors import BadRequestError

PDF_EXTENSIONS = {".pdf"}
DOCX_EXTENSIONS = {".docx"}
PPTX_EXTENSIONS = {".pptx"}
MARKDOWN_EXTENSIONS = {".md", ".markdown", ".mdown"}
HTML_EXTENSIONS = {".html", ".htm"}
TEXT_EXTENSIONS = {".txt", ".text", ".log"}

SUPPORTED_EXTENSIONS = (
    PDF_EXTENSIONS
    | DOCX_EXTENSIONS
    | PPTX_EXTENSIONS
    | MARKDOWN_EXTENSIONS
    | HTML_EXTENSIONS
    | TEXT_EXTENSIONS
)

MIME_TYPES = {
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    ".md": "text/markdown",
    ".markdown": "text/markdown",
    ".mdown": "text/markdown",
    ".html": "text/html",
    ".htm": "text/html",
    ".txt": "text/plain",
    ".text": "text/plain",
    ".log": "text/plain",
}

TEXT_DECODINGS = ("utf-8-sig", "utf-8", "gb18030", "big5", "latin-1")


class DocumentParseError(BadRequestError):
    code = "DOCUMENT_PARSE_FAILED"
    message = "无法解析该文件内容"


@dataclass(frozen=True)
class ParsedDocument:
    text: str
    source_type: SourceType
    metadata: dict[str, Any] = field(default_factory=dict)


def extension_of(filename: str) -> str:
    lowered = (filename or "").strip().lower()
    dot = lowered.rfind(".")
    return lowered[dot:] if dot >= 0 else ""


def is_supported_filename(filename: str) -> bool:
    return extension_of(filename) in SUPPORTED_EXTENSIONS


def mime_type_for(filename: str) -> str:
    return MIME_TYPES.get(extension_of(filename), "application/octet-stream")


def source_type_for(filename: str, declared: str | None = None) -> SourceType:
    if declared:
        try:
            return SourceType(declared.strip().lower())
        except ValueError:
            pass
    extension = extension_of(filename)
    if extension in PDF_EXTENSIONS:
        return SourceType.PDF
    if extension in DOCX_EXTENSIONS:
        return SourceType.DOCX
    if extension in PPTX_EXTENSIONS:
        return SourceType.PPTX
    if extension in MARKDOWN_EXTENSIONS:
        return SourceType.MARKDOWN
    if extension in HTML_EXTENSIONS:
        return SourceType.HTML
    if extension in TEXT_EXTENSIONS:
        return SourceType.TXT
    raise DocumentParseError("暂不支持该文件类型")


def decode_text(data: bytes) -> str:
    for encoding in TEXT_DECODINGS:
        try:
            return data.decode(encoding)
        except (UnicodeDecodeError, LookupError):
            continue
    return data.decode("utf-8", errors="replace")


class _HtmlTextExtractor(HTMLParser):
    SKIPPED_TAGS = {"script", "style", "noscript", "template", "head"}

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self._chunks: list[str] = []
        self._title_chunks: list[str] = []
        self._skip_depth = 0
        self._in_title = False

    def handle_starttag(self, tag: str, attrs) -> None:
        if tag in self.SKIPPED_TAGS:
            self._skip_depth += 1
        elif tag == "title":
            self._in_title = True
        elif tag in {"p", "div", "br", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6"}:
            self._chunks.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag in self.SKIPPED_TAGS and self._skip_depth > 0:
            self._skip_depth -= 1
        elif tag == "title":
            self._in_title = False

    def handle_data(self, data: str) -> None:
        if self._in_title:
            self._title_chunks.append(data)
            return
        if self._skip_depth == 0 and data.strip():
            self._chunks.append(data)

    @property
    def text(self) -> str:
        return _collapse_blank_lines("".join(self._chunks))

    @property
    def title(self) -> str | None:
        title = "".join(self._title_chunks).strip()
        return title or None


def _collapse_blank_lines(text: str) -> str:
    normalised = text.replace("\r\n", "\n").replace("\r", "\n")
    normalised = re.sub(r"[ \t\f\v]+", " ", normalised)
    normalised = re.sub(r"\n\s*\n\s*\n+", "\n\n", normalised)
    return normalised.strip()


def _parse_pdf(data: bytes) -> ParsedDocument:
    import pdfplumber

    pages: list[str] = []
    metadata: dict[str, Any] = {}
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        metadata["pageCount"] = len(pdf.pages)
        for page in pdf.pages:
            pages.append(page.extract_text() or "")
        info = pdf.metadata or {}
        for key in ("Title", "Author", "Subject", "Producer"):
            if info.get(key):
                metadata[key.lower()] = str(info[key])
    text = _collapse_blank_lines("\n\n".join(pages))
    if not text:
        raise DocumentParseError("该 PDF 没有可提取的文本内容")
    return ParsedDocument(text=text, source_type=SourceType.PDF, metadata=metadata)


def _parse_docx(data: bytes) -> ParsedDocument:
    import docx

    document = docx.Document(io.BytesIO(data))
    blocks = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            blocks.append("\t".join(cell.text for cell in row.cells))
    text = _collapse_blank_lines("\n".join(blocks))
    if not text:
        raise DocumentParseError("该 DOCX 没有可提取的文本内容")

    metadata: dict[str, Any] = {"paragraphCount": len(document.paragraphs)}
    properties = document.core_properties
    if properties.title:
        metadata["title"] = properties.title
    if properties.author:
        metadata["author"] = properties.author
    return ParsedDocument(text=text, source_type=SourceType.DOCX, metadata=metadata)


def _parse_pptx(data: bytes) -> ParsedDocument:
    from pptx import Presentation

    presentation = Presentation(io.BytesIO(data))
    blocks: list[str] = []
    for index, slide in enumerate(presentation.slides, start=1):
        blocks.append(f"## Slide {index}")
        for shape in slide.shapes:
            text_frame = getattr(shape, "text_frame", None)
            if text_frame is not None and text_frame.text.strip():
                blocks.append(text_frame.text)
    text = _collapse_blank_lines("\n".join(blocks))
    if not text:
        raise DocumentParseError("该 PPTX 没有可提取的文本内容")
    return ParsedDocument(
        text=text,
        source_type=SourceType.PPTX,
        metadata={"slideCount": len(presentation.slides)},
    )


def _parse_markdown(data: bytes) -> ParsedDocument:
    text = _collapse_blank_lines(decode_text(data))
    if not text:
        raise DocumentParseError("该 Markdown 文件内容为空")
    headings = re.findall(r"^#{1,6}\s+(.+)$", text, flags=re.MULTILINE)
    metadata: dict[str, Any] = {"headingCount": len(headings)}
    if headings:
        metadata["firstHeading"] = headings[0].strip()[:200]
    return ParsedDocument(text=text, source_type=SourceType.MARKDOWN, metadata=metadata)


def _parse_html(data: bytes) -> ParsedDocument:
    parser = _HtmlTextExtractor()
    parser.feed(decode_text(data))
    parser.close()
    text = parser.text
    if not text:
        raise DocumentParseError("该 HTML 文档没有可提取的文本内容")
    metadata: dict[str, Any] = {}
    if parser.title:
        metadata["title"] = parser.title
    return ParsedDocument(text=text, source_type=SourceType.HTML, metadata=metadata)


def _parse_text(data: bytes) -> ParsedDocument:
    text = _collapse_blank_lines(decode_text(data))
    if not text:
        raise DocumentParseError("该文本文件内容为空")
    return ParsedDocument(
        text=text,
        source_type=SourceType.TXT,
        metadata={"lineCount": text.count("\n") + 1},
    )


_NATIVE_PARSERS = {
    SourceType.PDF: _parse_pdf,
    SourceType.DOCX: _parse_docx,
    SourceType.PPTX: _parse_pptx,
    SourceType.MARKDOWN: _parse_markdown,
    SourceType.HTML: _parse_html,
    SourceType.TXT: _parse_text,
}


def _parse_with_markitdown(data: bytes) -> str | None:
    """Optional MarkItDown fallback. Returns ``None`` when unavailable."""

    try:
        from markitdown import MarkItDown
    except ImportError:
        return None
    try:
        result = MarkItDown().convert_stream(io.BytesIO(data))
    except Exception:  # pragma: no cover - fallback path
        return None
    return _collapse_blank_lines(result.text_content or "") or None


def parse_document(
    filename: str,
    data: bytes,
    *,
    declared_source_type: str | None = None,
) -> ParsedDocument:
    """Parse ``data`` into plain text plus format metadata."""

    if not data:
        raise DocumentParseError("文件内容为空")

    source_type = source_type_for(filename, declared_source_type)
    parser = _NATIVE_PARSERS.get(source_type)
    if parser is None:
        raise DocumentParseError("暂不支持该文件类型")

    try:
        return parser(data)
    except DocumentParseError:
        raise
    except Exception as exc:  # noqa: BLE001 - normalise third-party parser failures
        fallback = _parse_with_markitdown(data)
        if fallback:
            return ParsedDocument(
                text=fallback,
                source_type=source_type,
                metadata={"parser": "markitdown"},
            )
        raise DocumentParseError("文件可能已损坏或格式不受支持") from exc

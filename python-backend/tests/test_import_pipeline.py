"""Import pipeline: parsing, source assets, dedup and size/type guards."""

from __future__ import annotations

import io
import sqlite3

import pytest
from fastapi.testclient import TestClient

from knowledge_desk.config import Settings
from knowledge_desk.domain.models import KnowledgeItem
from knowledge_desk.main import create_app

MINIMAL_PDF = b"""%PDF-1.4
1 0 obj << /Type /Catalog /Pages 2 0 R >> endobj
2 0 obj << /Type /Pages /Kids [3 0 R] /Count 1 >> endobj
3 0 obj << /Type /Page /Parent 2 0 R /MediaBox [0 0 300 300] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >> endobj
4 0 obj << /Length 64 >> stream
BT /F1 12 Tf 20 250 Td (Knowledge Desk PDF sample) Tj ET
endstream endobj
5 0 obj << /Type /Font /Subtype /Type1 /BaseFont /Helvetica >> endobj
trailer << /Root 1 0 R /Size 6 >>
%%EOF
"""


def _docx_bytes(paragraphs: list[str]) -> bytes:
    import docx

    document = docx.Document()
    for text in paragraphs:
        document.add_paragraph(text)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _pptx_bytes(title: str, body: str) -> bytes:
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = title
    box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(1))
    box.text_frame.text = body
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def _upload(client: TestClient, auth: dict[str, str], filename: str, data: bytes, title: str | None = None):
    files = {"file": (filename, data, "application/octet-stream")}
    form = {"title": title} if title else None
    return client.post(
        "/api/v1/knowledge-items/import/upload", files=files, data=form, headers=auth
    )


@pytest.mark.parametrize(
    ("filename", "payload", "expected_type", "needle"),
    (
        ("notes.md", "# 标题\n\n这是 Markdown 正文，包含足够多的词语用于计数。".encode(), "markdown", "Markdown"),
        ("notes.txt", "纯文本内容，用来验证 TXT 解析路径。".encode(), "txt", "TXT"),
        ("page.html", b"<html><head><title>HTML \xe6\xa0\x87\xe9\xa2\x98</title></head><body><p>\xe6\xad\xa3\xe6\x96\x87\xe5\x86\x85\xe5\xae\xb9</p></body></html>", "html", "HTML"),
    ),
)
def test_text_formats_parse(
    client: TestClient, auth: dict[str, str], filename: str, payload: bytes, expected_type: str, needle: str
) -> None:
    response = _upload(client, auth, filename, payload, title=f"{needle} 资料")
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["sourceType"] == expected_type
    assert body["wordCount"] > 0
    assert body["sourceAsset"]["originalFilename"] == filename
    assert body["sourceAsset"]["availability"] == "available"
    assert body["sourceAsset"]["origin"] == "picker"


def test_pdf_upload_extracts_text(client: TestClient, auth: dict[str, str]) -> None:
    response = _upload(client, auth, "sample.pdf", MINIMAL_PDF, title="PDF 资料")
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["sourceType"] == "pdf"
    assert "Knowledge Desk PDF sample" in body["rawContent"]
    assert body["sourceAsset"]["mediaType"] == "application/pdf"


def test_docx_upload_extracts_paragraphs(client: TestClient, auth: dict[str, str]) -> None:
    data = _docx_bytes(["知识工作台文档", "第二段用于验证段落抽取。"])
    response = _upload(client, auth, "doc.docx", data)
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["sourceType"] == "docx"
    assert "知识工作台文档" in body["rawContent"]
    assert "第二段用于验证段落抽取。" in body["rawContent"]
    # Title falls back to the derived first line when none is supplied.
    assert body["title"] == "知识工作台文档"


def test_pptx_upload_extracts_slide_text(client: TestClient, auth: dict[str, str]) -> None:
    data = _pptx_bytes("幻灯片标题", "幻灯片正文内容")
    response = _upload(client, auth, "deck.pptx", data, title="PPT 资料")
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["sourceType"] == "pptx"
    assert "幻灯片标题" in body["rawContent"]
    assert "幻灯片正文内容" in body["rawContent"]


def test_uploaded_asset_is_stored_under_data_dir(
    client: TestClient, auth: dict[str, str], settings: Settings, session_factory
) -> None:
    response = _upload(client, auth, "stored.md", "# 受管原件\n\n正文内容。".encode())
    assert response.status_code == 200
    item_id = response.json()["id"]

    session = session_factory()
    try:
        item = session.get(KnowledgeItem, item_id)
        asset = item.source_asset
        assert asset is not None
        # Stored path is relative to the managed sources directory.
        assert not asset.stored_relative_path.startswith("/")
        stored = settings.sources_dir / asset.stored_relative_path
        assert stored.exists()
        assert stored.is_relative_to(settings.sources_dir)
        # The API response must not leak the managed path.
        assert str(settings.data_dir) not in response.text
    finally:
        session.close()


def test_response_never_leaks_content_hash(client: TestClient, auth: dict[str, str], session_factory) -> None:
    response = _upload(client, auth, "hash.md", "# 哈希\n\n正文。".encode())
    item_id = response.json()["id"]

    session = session_factory()
    try:
        digest = session.get(KnowledgeItem, item_id).content_hash
    finally:
        session.close()

    assert digest and len(digest) == 64
    assert digest not in response.text


def test_duplicate_upload_conflicts(client: TestClient, auth: dict[str, str]) -> None:
    data = _docx_bytes(["重复文档", "内容一致。"])
    first = _upload(client, auth, "dup.docx", data)
    assert first.status_code == 200

    second = _upload(client, auth, "dup.docx", data)
    assert second.status_code == 409
    payload = second.json()
    assert payload["code"] == "DUPLICATE_CONTENT"
    assert "already" in payload["message"].lower()


def test_duplicate_file_payload_conflicts(client: TestClient, auth: dict[str, str]) -> None:
    payload = {"content": "完全一致的本地文件内容。", "sourceType": "markdown"}
    assert client.post(
        "/api/v1/knowledge-items/import/file", json=payload, headers=auth
    ).status_code == 200
    conflict = client.post("/api/v1/knowledge-items/import/file", json=payload, headers=auth)
    assert conflict.status_code == 409
    assert conflict.json()["code"] == "DUPLICATE_CONTENT"


def test_snippet_repeat_import_is_allowed(client: TestClient, auth: dict[str, str]) -> None:
    """Non-file captures are intentionally not deduplicated."""

    payload = {"content": "同一段手动片段可以重复导入。"}
    assert client.post(
        "/api/v1/knowledge-items/import/snippet", json=payload, headers=auth
    ).status_code == 200
    assert client.post(
        "/api/v1/knowledge-items/import/snippet", json=payload, headers=auth
    ).status_code == 200


def test_preflight_reports_known_hashes(client: TestClient, auth: dict[str, str]) -> None:
    data = _docx_bytes(["预检文档", "内容。"])
    assert _upload(client, auth, "pre.docx", data).status_code == 200

    import hashlib

    digest = hashlib.sha256(data).hexdigest()
    other = hashlib.sha256(b"absent").hexdigest()

    response = client.post(
        "/api/v1/knowledge-items/import/preflight",
        json={"contentHashes": [digest, other]},
        headers=auth,
    )
    assert response.status_code == 200
    body = response.json()
    assert body["existingHashes"] == [digest]
    assert body["knownCount"] == 1

    invalid = client.post(
        "/api/v1/knowledge-items/import/preflight", json={"contentHashes": ["not-a-hash"]}, headers=auth
    )
    assert invalid.status_code == 400

    empty = client.post(
        "/api/v1/knowledge-items/import/preflight", json={"contentHashes": []}, headers=auth
    )
    assert empty.status_code == 400


def test_unsupported_file_type_is_rejected(client: TestClient, auth: dict[str, str]) -> None:
    response = _upload(client, auth, "archive.zip", b"PK\x03\x04binary")
    assert response.status_code == 415
    assert response.json()["code"] == "UNSUPPORTED_MEDIA_TYPE"


def test_corrupted_document_is_rejected(client: TestClient, auth: dict[str, str]) -> None:
    response = _upload(client, auth, "broken.docx", b"this is not a real docx")
    assert response.status_code == 400
    assert response.json()["code"] == "DOCUMENT_PARSE_FAILED"
    assert "/Users" not in response.text


def test_oversize_upload_is_rejected(tmp_path, auth_factory) -> None:
    settings = Settings(
        data_dir=tmp_path / "small",
        jwt_secret="oversize-upload-jwt-secret-tests-only",
        db_encryption_key="oversize-key",
        desktop_mode=True,
        log_level="WARNING",
        max_upload_bytes=2048,
    )
    app = create_app(settings)
    with TestClient(app) as client:
        auth = auth_factory(client)
        response = _upload(client, auth, "big.txt", b"x" * 4096)
        assert response.status_code == 413
        assert response.json()["code"] == "PAYLOAD_TOO_LARGE"


def test_sqlite_records_item_without_absolute_paths(settings: Settings, auth: dict[str, str], client: TestClient) -> None:
    _upload(client, auth, "path.md", "# 路径\n\n内容。".encode())
    connection = sqlite3.connect(settings.database_path)
    try:
        rows = connection.execute("select stored_relative_path from source_assets").fetchall()
    finally:
        connection.close()
    assert rows
    for (relative_path,) in rows:
        assert not relative_path.startswith("/")
        assert "Users" not in relative_path

"""Unit coverage for security primitives, redaction, text helpers and parsing."""

from __future__ import annotations

import io
from datetime import datetime, timedelta, timezone

import pytest

from knowledge_desk.application.text_utils import (
    count_words,
    derive_title,
    detect_language,
    heuristic_organize,
    normalize_text,
)
from knowledge_desk.errors import EncryptionUnavailableError, UnauthorizedError
from knowledge_desk.infrastructure.crypto import SecretCipher
from knowledge_desk.infrastructure.document_parser import (
    DocumentParseError,
    extension_of,
    is_supported_filename,
    mime_type_for,
    parse_document,
    source_type_for,
)
from knowledge_desk.infrastructure.redaction import redact, redact_max
from knowledge_desk.infrastructure.security import (
    create_access_token,
    create_refresh_token,
    decode_token,
    hash_password,
    verify_password,
)
from knowledge_desk.main import RedactingFilter
from knowledge_desk.domain.enums import SourceType


# --------------------------------------------------------------------------- #
# Passwords and tokens
# --------------------------------------------------------------------------- #


def test_password_hash_roundtrip_is_salted() -> None:
    first = hash_password("Desktop!1234Aa")
    second = hash_password("Desktop!1234Aa")

    assert first.startswith("scrypt$")
    assert first != second  # unique salt per hash
    assert verify_password("Desktop!1234Aa", first)
    assert not verify_password("Desktop!1234ab", first)


def test_password_verification_rejects_malformed_hashes() -> None:
    for broken in ("", "plaintext", "scrypt$bad", "bcrypt$1$2$3$4$5"):
        assert verify_password("Desktop!1234Aa", broken) is False


def test_access_token_roundtrip_and_expiry() -> None:
    secret = "unit-test-secret-long-enough-for-hs256"
    now = datetime.now(timezone.utc)

    token = create_access_token("user-1", secret, 60, now=now)
    payload = decode_token(token, secret, expected_type="access")
    assert payload["sub"] == "user-1"
    assert payload["typ"] == "access"

    expired = create_access_token("user-1", secret, 1, now=now - timedelta(days=1))
    with pytest.raises(UnauthorizedError):
        decode_token(expired, secret)


def test_token_type_confusion_is_rejected() -> None:
    refresh = create_refresh_token("user-1", "unit-test-secret-long-enough-for-hs256", 1)
    with pytest.raises(UnauthorizedError):
        decode_token(refresh, "unit-test-secret-long-enough-for-hs256", expected_type="access")


def test_token_signed_with_other_secret_is_rejected() -> None:
    token = create_access_token("user-1", "unit-test-secret-value-number-one", 60)
    with pytest.raises(UnauthorizedError):
        decode_token(token, "unit-test-secret-value-number-two")


# --------------------------------------------------------------------------- #
# Credential encryption
# --------------------------------------------------------------------------- #


def test_cipher_encrypts_and_masks() -> None:
    cipher = SecretCipher("runtime-key-material")
    assert cipher.available

    encrypted = cipher.encrypt("sk-abcdefghijklmnop1234")
    assert encrypted.startswith("enc:v1:")
    assert "sk-abcdefghijklmnop1234" not in encrypted
    assert cipher.decrypt(encrypted) == "sk-abcdefghijklmnop1234"
    assert SecretCipher.mask("sk-abcdefghijklmnop1234") == "sk-…1234"
    assert SecretCipher.mask("short") == "••••"
    assert SecretCipher.mask(None) is None


def test_cipher_without_key_refuses_to_encrypt() -> None:
    cipher = SecretCipher(None)
    assert not cipher.available
    assert cipher.decrypt("enc:v1:whatever") is None
    with pytest.raises(EncryptionUnavailableError):
        cipher.encrypt("sk-value")


def test_cipher_cannot_decrypt_with_another_key() -> None:
    encrypted = SecretCipher("key-a").encrypt("sk-value")
    assert SecretCipher("key-b").decrypt(encrypted) is None


# --------------------------------------------------------------------------- #
# Redaction
# --------------------------------------------------------------------------- #


def test_redact_masks_credentials_tokens_and_home_paths() -> None:
    sample = (
        "key=sk-abcdefghijklmnop1234 Authorization: Bearer eyJhbGciOiJIUzI1NiJ9.abcdefgh.ijklmnop "
        "reading /Users/someone/Documents/private/note.md"
    )
    cleaned = redact(sample)
    assert "sk-abcdefghijklmnop1234" not in cleaned
    assert "eyJhbGciOiJIUzI1NiJ9" not in cleaned
    assert "/Users/someone" not in cleaned
    assert "[redacted]" in cleaned


def test_redact_max_clamps_length() -> None:
    assert len(redact_max("x" * 500)) == 255
    assert redact_max(None) is None


def test_logging_filter_redacts_records() -> None:
    import logging

    filter_ = RedactingFilter()
    record = logging.LogRecord(
        "test", logging.INFO, __file__, 1, "using key sk-abcdefghijklmnop1234", None, None
    )
    assert filter_.filter(record) is True
    assert "sk-abcdefghijklmnop1234" not in record.getMessage()


# --------------------------------------------------------------------------- #
# Text helpers
# --------------------------------------------------------------------------- #


def test_normalize_and_word_count() -> None:
    raw = "  第一行  \r\n\r\n\r\n  第二行   \n第二行   \n<p>标签</p>"
    normalized = normalize_text(raw)
    assert "\r" not in normalized
    assert "\n\n\n" not in normalized
    assert "<p>" not in normalized
    assert count_words("hello world 你好世界") == 2 + 4
    assert count_words("") == 0


def test_language_detection() -> None:
    assert detect_language("这是一段中文内容，用于语言检测。") == "zh"
    assert detect_language("This is a plain English sentence used for detection.") == "en"
    assert detect_language("") == "und"


def test_derive_title_prefers_explicit_then_heading() -> None:
    assert derive_title("显式标题", "正文", "fallback") == "显式标题"
    assert derive_title(None, "# Markdown 标题\n\n正文", "fallback") == "Markdown 标题"
    # Lines shorter than four characters are skipped until one qualifies.
    assert derive_title(None, "短\n\n较长的一行正文内容", "fallback") == "较长的一行正文内容"
    assert derive_title(None, "短\n\n正文", "fallback") == "短"
    assert derive_title(None, "", "本地文件") == "本地文件"


def test_heuristic_organize_is_deterministic() -> None:
    text = (
        "FastAPI 是 Python 的现代 Web 框架。\n\n"
        "FastAPI 使用 Pydantic 做数据校验，FastAPI 也支持依赖注入。\n\n"
        "部署时通常搭配 Uvicorn。"
    )
    first = heuristic_organize("FastAPI 入门", text)
    second = heuristic_organize("FastAPI 入门", text)

    assert first == second
    assert first.summary.startswith("FastAPI 是 Python 的现代 Web 框架。")
    assert first.cleaned_content
    assert "fastapi" in first.tags
    assert len(first.tags) <= 5


def test_heuristic_summary_is_truncated() -> None:
    long_text = "字" * 400
    draft = heuristic_organize("长文", long_text)
    assert draft.summary.endswith("…")
    assert len(draft.summary) <= 241


# --------------------------------------------------------------------------- #
# Document parsing
# --------------------------------------------------------------------------- #


def test_extension_and_mime_mapping() -> None:
    assert extension_of("A.PDF") == ".pdf"
    assert mime_type_for("a.md") == "text/markdown"
    assert mime_type_for("a.docx").endswith("wordprocessingml.document")
    assert mime_type_for("unknown.bin") == "application/octet-stream"
    assert is_supported_filename("a.pptx")
    assert not is_supported_filename("a.exe")


def test_source_type_resolution() -> None:
    assert source_type_for("a.pdf") is SourceType.PDF
    assert source_type_for("a.docx") is SourceType.DOCX
    assert source_type_for("a.pptx") is SourceType.PPTX
    assert source_type_for("a.md") is SourceType.MARKDOWN
    assert source_type_for("a.html") is SourceType.HTML
    assert source_type_for("a.txt") is SourceType.TXT
    assert source_type_for("a.dat", declared="snippet") is SourceType.SNIPPET
    with pytest.raises(DocumentParseError):
        source_type_for("a.xyz")


def test_html_extraction_uses_title_and_drops_scripts() -> None:
    html = (
        "<html><head><title>页面标题</title><script>var a = 1;</script></head>"
        "<body><p>第一段正文。</p><p>第二段正文。</p></body></html>"
    )
    parsed = parse_document("page.html", html.encode())
    assert parsed.metadata["title"] == "页面标题"
    assert "第一段正文。" in parsed.text
    assert "var a = 1" not in parsed.text


def test_gb18030_text_is_decoded() -> None:
    parsed = parse_document("legacy.txt", "中文编码测试内容。".encode("gb18030"))
    assert "中文编码测试内容。" in parsed.text


def test_empty_and_corrupt_documents_are_rejected() -> None:
    with pytest.raises(DocumentParseError):
        parse_document("empty.txt", b"")
    with pytest.raises(DocumentParseError):
        parse_document("empty.txt", b"   \n  ")
    with pytest.raises(DocumentParseError):
        parse_document("broken.docx", b"not a docx at all")


def test_unicode_filename_is_preserved() -> None:
    parsed = parse_document("中文资料.md", "# 标题\n\n正文内容。".encode())
    assert parsed.source_type is SourceType.MARKDOWN
    assert parsed.metadata["firstHeading"] == "标题"


def test_docx_and_pptx_metadata_are_reported() -> None:
    import docx
    from pptx import Presentation
    from pptx.util import Inches

    document = docx.Document()
    document.add_paragraph("段落一")
    docx_buffer = io.BytesIO()
    document.save(docx_buffer)
    docx_parsed = parse_document("memo.docx", docx_buffer.getvalue())
    assert docx_parsed.metadata["paragraphCount"] == 1
    assert "段落一" in docx_parsed.text

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "标题"
    box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(4), Inches(1))
    box.text_frame.text = "内容"
    pptx_buffer = io.BytesIO()
    presentation.save(pptx_buffer)
    pptx_parsed = parse_document("deck.pptx", pptx_buffer.getvalue())
    assert pptx_parsed.metadata["slideCount"] == 1
    assert "标题" in pptx_parsed.text

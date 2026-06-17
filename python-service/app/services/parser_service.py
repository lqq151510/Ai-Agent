"""文档解析服务。

优先使用 MarkItDown 进行通用转换；对于需要更精细结构化信息的格式
（pdf / docx / pptx），补充使用 pdfplumber、python-docx、python-pptx
提取元数据（页数、段落数、幻灯片数等），保证返回内容更完整。
"""

import logging
import os
from typing import Dict, Tuple

from markitdown import MarkItDown

logger = logging.getLogger(__name__)
md = MarkItDown()

# 支持的文件格式（小写扩展名 -> 描述）
SUPPORTED_FORMATS = {
    ".pdf": "pdf",
    ".docx": "docx",
    ".doc": "docx",
    ".pptx": "pptx",
    ".ppt": "pptx",
    ".md": "md",
    ".markdown": "md",
    ".txt": "txt",
    ".html": "html",
    ".htm": "html",
}


def get_source_format(file_path: str) -> str:
    """根据扩展名返回归一化的源格式标识。"""
    ext = os.path.splitext(file_path)[1].lower()
    return SUPPORTED_FORMATS.get(ext, ext.lstrip("."))


def convert_with_markitdown(file_path: str) -> str:
    """使用 MarkItDown 执行阻塞式转换，返回 Markdown 文本。"""
    result = md.convert(file_path)
    return result.text_content


def parse_document(file_path: str) -> Tuple[str, Dict[str, str]]:
    """解析文档，返回 (markdown, metadata)。

    对于结构化文档（pdf/docx/pptx），在 MarkItDown 转换基础上补充元数据。
    对于 md/txt/html，直接使用 MarkItDown 转换结果。
    """
    source_format = get_source_format(file_path)
    metadata: Dict[str, str] = {"sourceFormat": source_format}

    # 先用 MarkItDown 拿到主体 Markdown
    markdown = convert_with_markitdown(file_path)

    # 补充各格式的元数据
    try:
        if source_format == "pdf":
            metadata.update(_extract_pdf_metadata(file_path))
        elif source_format == "docx":
            metadata.update(_extract_docx_metadata(file_path))
        elif source_format == "pptx":
            metadata.update(_extract_pptx_metadata(file_path))
    except Exception as e:  # 元数据提取失败不影响主流程
        logger.warning("提取 %s 元数据失败: %s", source_format, e)

    return markdown, metadata


def _extract_pdf_metadata(file_path: str) -> Dict[str, str]:
    """使用 pdfplumber 提取 PDF 元数据。"""
    import pdfplumber

    meta: Dict[str, str] = {}
    with pdfplumber.open(file_path) as pdf:
        meta["pageCount"] = str(len(pdf.pages))
        info = pdf.metadata or {}
        if info.get("Title"):
            meta["title"] = str(info["Title"])
        if info.get("Author"):
            meta["author"] = str(info["Author"])
    return meta


def _extract_docx_metadata(file_path: str) -> Dict[str, str]:
    """使用 python-docx 提取 DOCX 元数据。"""
    import docx

    meta: Dict[str, str] = {}
    doc = docx.Document(file_path)
    meta["paragraphCount"] = str(len(doc.paragraphs))
    core = doc.core_properties
    if core.title:
        meta["title"] = core.title
    if core.author:
        meta["author"] = core.author
    return meta


def _extract_pptx_metadata(file_path: str) -> Dict[str, str]:
    """使用 python-pptx 提取 PPTX 元数据。"""
    from pptx import Presentation

    meta: Dict[str, str] = {}
    prs = Presentation(file_path)
    meta["slideCount"] = str(len(prs.slides))
    core = prs.core_properties
    if core.title:
        meta["title"] = core.title
    if core.author:
        meta["author"] = core.author
    return meta
